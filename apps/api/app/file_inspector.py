"""Safe, deterministic inspection of verification output files."""
from __future__ import annotations

import csv
import hashlib
import io
import json
from pathlib import Path
from typing import Any


def inspect_file(filename: str, content: bytes) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    result: dict[str, Any] = {"filename": filename, "format": suffix.lstrip("."), "bytes": len(content), "sha256": hashlib.sha256(content).hexdigest(), "valid": True}
    if suffix == ".csv":
        rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        result.update({"columns": list(rows[0].keys()) if rows else [], "row_count": len(rows), "sample": rows[:3]})
    elif suffix == ".json":
        value = json.loads(content.decode("utf-8"))
        rows = value if isinstance(value, list) else [value]
        result.update({"columns": list(rows[0].keys()) if rows and isinstance(rows[0], dict) else [], "row_count": len(rows), "sample": rows[:3]})
    elif suffix in {".txt", ".md"}:
        text = content.decode("utf-8", errors="replace")
        result.update({"characters": len(text), "line_count": len(text.splitlines()), "preview": text[:500]})
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheets: dict[str, Any] = {}
            for sheet in workbook.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                headers = [str(v) for v in rows[0]] if rows else []
                sheets[sheet.title] = {"columns": headers, "row_count": max(0, len(rows) - 1), "sample": [dict(zip(headers, row)) for row in rows[1:4]]}
            result["sheets"] = sheets
        except ImportError:
            result.update({"valid": False, "error": "XLSX inspection requires openpyxl"})
    else:
        result.update({"valid": False, "error": f"Unsupported file format: {suffix or 'unknown'}"})
    return result
